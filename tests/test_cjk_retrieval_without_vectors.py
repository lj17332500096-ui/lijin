"""护栏：中文关键词检索**不得**依赖 FTS5 的分词器。

⚠️ **为什么需要这条**（10-05 CI run 37326450894 实据）：
SQLite FTS5 的默认 `unicode61` 分词器把**中文整段当成一个 token**。
内容「按钮必须放在右下角且带确认弹窗。」用 `"按钮"` / `"确认弹窗"` / `"右下角"`
查，命中数一律为 **0**。这是**分词器问题，不是查询写法问题** ——
把查询改成二元组 OR（`"按钮" OR "钮必" OR ...`）实测同样 0 hits。

`sources/retriever.py` 曾只走 FTS5 那条路 ⇒ 中文检索恒失效。本地全绿是因为
`rag.py` 的 ONNX 向量把结果兜住了；而 `requirements.txt` 无 onnxruntime、
模型目录 `models/bge-small-zh-v1.5` 也不入库 ⇒ **CI 上向量必然不可用**
⇒ 掩盖消失，暴露成 `mode='no_results'`、检索 0 条。

判据设计（双向）：
  ① **无向量时中文必须能召回**（模拟 CI）—— 这是真正要防的回归；
  ② **无关中文不得被误召回** —— 防「用 or 匹配一切」这种假修复；
  ③ 直接钉住 FTS5 分词器行为，让「换分词器 / 改查询写法」这类无效修法
     在测试里就能看见，而不是等 CI。
"""
from __future__ import annotations

import math
import re
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
if str(BASE) not in sys.path:
    sys.path.insert(0, str(BASE))

#: 内容里出现的检索目标
_DOC = "按钮必须放在右下角且带确认弹窗。"
#: 与 _DOC 中确实共现的 CJK 查询
_QUERIES = ("按钮", "右下角", "确认弹窗")
#: 与 _DOC 毫无 CJK 字符交集的查询 —— 用于误伤方向
_UNRELATED = ("量子力学入门书籍推荐", "红烧肉的做法", "天文望远镜选购")


def _fts5_shortname_break(content: str, terms: tuple[str, ...]) -> int:
    """在最小 FTS5 库里复刻索引侧写入 + 查询侧phrase 匹配，返回命中数。"""
    db = Path(tempfile.mkdtemp()) / "fts_probe.db"
    conn = sqlite3.connect(str(db))
    try:
        conn.execute("CREATE VIRTUAL TABLE f USING fts5(content)")
        conn.execute("INSERT INTO f(rowid, content) VALUES (?,?)", (1, content))
        expr = " OR ".join('"%s"' % t for t in terms)
        return conn.execute("SELECT count(*) FROM f WHERE f MATCH ?", (expr,)).fetchone()[0]
    finally:
        conn.close()


class Fts5ChineseTokenizerTests(unittest.TestCase):
    """钉住「FTS5 unicode61 对中文整段分词」这个事实本身。

    这不是缺陷，是**前提**。若哪天它变了（换分词器 / 换 SQLite 版本），
    本组会红，提示可以简化 `sources/retriever.py` 的 Python 侧兜底。
    """

    def test_cjk_phrase_never_matches_across_token_boundary(self) -> None:
        for term in _QUERIES:
            with self.subTest(term=term):
                hits = _fts5_shortname_break(_DOC, (term,))
                self.assertEqual(
                    hits, 0,
                    f"FTS5 对「{term}」命中 {hits} 条 —— unicode61 的中文行为已变。"
                    "这会让 sources/retriever.py 的 Python 侧兜底变成冗余，"
                    "请重新评估是否可简化（**先有实据再改**）。",
                )

    def test_bigram_ors_also_cannot_rescue(self) -> None:
        """**排除一种无效修法**：以为「查询侧改二元组」就够了。实测 0 hits。"""
        bigrams = tuple("".join(p) for p in zip(_DOC, _DOC[1:]))
        cjk_only = tuple(b for b in bigrams if re.fullmatch(r"[\u4e00-\u9fff]{2}", b))
        self.assertTrue(cjk_only, "前提失效：没切出中文二元组")
        self.assertEqual(_fts5_shortname_break(_DOC, cjk_only), 0)


class ChineseLexicalFallbackTests(unittest.TestCase):
    """sources 侧的中文召回必须**不依赖向量**（CI 上向量必然不可用）。"""

    def test_cjk_bigrams_cover_terms(self) -> None:
        from sources.retriever import cjk_bigrams

        for term in ("确认弹窗", "右下角"):
            with self.subTest(term=term):
                bigrams = cjk_bigrams(term)
                # term 内部相邻二元组都应产出
                self.assertEqual(
                    bigrams, ["".join(p) for p in zip(term, term[1:])],
                )

    def test_cjk_bigrams_split_across_text_boundary(self) -> None:
        """相邻二元组要跨词边界产出（这样「右下角」能命中「放在右下角」）。"""
        from sources.retriever import cjk_bigrams

        self.assertIn("右下", cjk_bigrams("放在右下角"))

    def test_match_expr_never_sends_whole_cjk_run_to_fts(self) -> None:
        """CJK 整段绝不能进 FTS 表达式（那是恒 0 命中的写法）。"""
        from sources.retriever import build_match_expr

        expr = build_match_expr("按钮 确认弹窗 右下角")
        self.assertNotIn('"确认弹窗"', expr)
        self.assertIn('"按钮"', expr)          # 二元组仍在
        self.assertIn('"确认"', expr)

    def test_python_side_scores_chinese_without_vectors(self) -> None:
        """端到端：关掉向量后中文仍能召回（这才是 CI 的真实条件）。"""
        import rag as rag_mod
        from sources import retriever
        from sources.indexer import index_source_file
        import tests.test_sources_rag as T

        orig = (rag_mod._EMBEDDER, rag_mod._EMBED_ERROR, rag_mod._EMBEDDER_LOADED)
        # 模拟 CI：向量彻底不可用
        rag_mod._EMBEDDER, rag_mod._EMBED_ERROR, rag_mod._EMBEDDER_LOADED = (
            None, "forced_off_for_test", True)
        tc = T.SourcesRagTests("test_docx_pdf_sources")
        tc.setUp()
        try:
            import docx as _docx

            pid = tc.container_a["id"]
            d = tc.src_dir_a / "spec.docx"
            doc = _docx.Document()
            doc.add_heading("界面规范", level=1)
            doc.add_paragraph(_DOC)
            doc.save(str(d))
            row = tc.manager.add_project_source(
                task_id=pid, display_name="spec.docx", stored_path=str(d),
                mime_type="application/vnd.openxmlformats-officedocument."
                          "wordprocessingml.document",
                size_bytes=d.stat().st_size, sha256="d" * 64,
                parse_status="pending")
            # 索引成功是前提 —— 否则下面的召回断言测的是别的东西
            self.assertTrue(index_source_file(pid, row).get("ok"))

            # ① 召回方向：中文必须能查到
            for q in _QUERIES:
                with self.subTest(query=q):
                    self.assertGreaterEqual(
                        retriever.retrieve(pid, q)["count"], 1,
                        f"无向量时中文「{q}」召回 0 条 —— "
                        "FTS5 中文失效的兜底没生效（或被改回去了）")

            # ② 误伤方向：无关中文不得被召回
            for q in _UNRELATED:
                with self.subTest(query=q):
                    self.assertEqual(
                        retriever.retrieve(pid, q)["count"], 0,
                        f"无关查询「{q}」被召回 ⇒ 判据已空转（什么都返回）")
        finally:
            tc.tearDown()
            (rag_mod._EMBEDDER, rag_mod._EMBED_ERROR,
             rag_mod._EMBEDDER_LOADED) = orig


class GuardSelfCheckTests(unittest.TestCase):
    """反空转：本护栏自身不能空转。"""

    def test_unrelated_queries_actually_have_no_cjk_overlap(self) -> None:
        """误伤用例必须与 _DOC 无 CJK 字符交集，否则「不召回」是废话。"""
        doc_chars = set(re.findall(r"[\u4e00-\u9fff]", _DOC))
        for q in _UNRELATED:
            with self.subTest(query=q):
                self.assertEqual(set(re.findall(r"[\u4e00-\u9fff]", q)) & doc_chars,
                                 set(), f"误伤用例「{q}」与文档有共同汉字，判据无效")

    def test_related_queries_actually_overlap(self) -> None:
        """召回用例必须与 _DOC 有 CJK 交集，否则「能召回」是废话。"""
        doc_chars = set(re.findall(r"[\u4e00-\u9fff]", _DOC))
        for q in _QUERIES:
            with self.subTest(query=q):
                self.assertTrue(set(re.findall(r"[\u4e00-\u9fff]", q)) & doc_chars)

    def test_fts_probe_returns_int_not_bogus(self) -> None:
        """探针本身要能工作：英文查询在同一个库里必须能命中。"""
        # 若这���都拿不到非 0，说明探针/表达式构造坏了，本护栏整体无鉴别力
        self.assertEqual(_fts5_shortname_break("hello world", ("hello",)), 1)
        self.assertEqual(_fts5_shortname_break("hello world", ("absent",)), 0)


if __name__ == "__main__":
    unittest.main()